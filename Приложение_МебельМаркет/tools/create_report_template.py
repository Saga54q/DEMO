"""Создаёт DOCX-каркас отчёта; серые зоны нужно заменить реальными скриншотами."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Отчёт_скриншоты_ШАБЛОН.docx"

def shade(cell, color):
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(element)

def set_font(run, name="Arial", size=10, color=None, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if bold is not None:
        run.bold = bold

def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.65)
    section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.7)
    section.right_margin = Inches(.7)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ОТЧЁТ О ПРОВЕРКЕ ПРИЛОЖЕНИЯ\n")
    set_font(run, size=20, color=(0, 100, 100), bold=True)
    run = p.add_run("МебельМаркет · вариант № 5 · В5_КОД 09.02.07-2-2026-ПА")
    set_font(run, size=12, bold=True)
    doc.add_paragraph("Шаблон. В серые области вставить реальные скриншоты окон приложения после запуска на рабочем ПК. Макеты и заглушки не считаются скриншотами.")
    doc.add_heading("Автоматические проверки", level=1)
    doc.add_paragraph("Запуск: python -m unittest -v. После выполнения тестов внесите фактический результат и дату проверки.")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ("Проверка", "Ожидаемый результат", "Результат / дата")):
        cell.text = text
    rows = [
        ("Импорт и схема", "Таблицы заполнены; PRAGMA foreign_key_check пуст", "________________"),
        ("Роли и вход", "Гость/клиент без инструментов; менеджер и администратор по правам", "________________"),
        ("Каталог", "Все поля, складские остатки, поиск/фильтры/сортировка по роли", "________________"),
        ("Скидка и остаток", ">15% teal #008080; остаток 0 — серый; цены оформлены по условию", "________________"),
        ("Товары и заказы", "CRUD администратора; товар из заказа удалить нельзя", "________________"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text

    cases = [
        ("1. Окно входа и гостевой режим", "Показать экран входа, вход гостем и список без фильтров/сортировки/поиска."),
        ("2. Клиент и ФИО", "Показать вход авторизованного клиента, ФИО/роль в шапке и просмотр каталога."),
        ("3. Склад и карточка товара", "Показать фото или picture.png, все атрибуты товара и остатки по складам."),
        ("4. Скидка и нулевой остаток", "Показать мебель со скидкой >15% в #008080, исходную красную цену с зачёркиванием, итоговую чёрную цену и нулевой остаток серым."),
        ("5. Менеджер: поиск, фильтр, сортировка и заказы", "Показать инструменты менеджера и просмотр списка заказов."),
        ("6. Администратор: товары", "Показать добавление/изменение товара и складского остатка."),
        ("7. Администратор: заказы", "Показать добавление/изменение/удаление заказа и запрет удаления товара, который заказан."),
        ("8. Выход", "Показать возврат к окну входа после выхода из учётной записи."),
    ]
    for title, note in cases:
        doc.add_page_break()
        doc.add_heading(title, level=1)
        doc.add_paragraph(note)
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = "ВСТАВИТЬ РЕАЛЬНЫЙ СКРИНШОТ ПРИЛОЖЕНИЯ\nWin + Shift + S → выделить окно → Ctrl + V"
        shade(cell, "E8ECEF")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.height = Inches(4.8)
        doc.add_paragraph("Подпись: ______________________________________________________________")
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == "__main__":
    main()
