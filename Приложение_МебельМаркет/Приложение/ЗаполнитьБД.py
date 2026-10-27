"""Импортирует условный демонстрационный Excel-набор в пустую Access-БД."""
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path
import sys

import pyodbc
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "БД" / "МебельМаркет.accdb"
XLSX_PATH = ROOT / "БД" / "Демо_импорт.xlsx"
DRIVER = "Microsoft Access Driver (*.mdb, *.accdb)"


def connect():
    if not DB_PATH.exists():
        raise FileNotFoundError(f"Не найдена база Access: {DB_PATH}\nСначала создайте её по инструкции.")
    available = pyodbc.drivers()
    if DRIVER not in available:
        raise RuntimeError("Не найден ODBC-драйвер Microsoft Access. Установите Access Database Engine той же разрядности, что и Python.")
    return pyodbc.connect(f"DRIVER={{{DRIVER}}};DBQ={DB_PATH};", autocommit=False)


def records(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value) for value in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]


def password_hash(value):
    return sha256(str(value).encode("utf-8")).hexdigest()


def main():
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"Нет файла демонстрационных данных: {XLSX_PATH}")
    book = load_workbook(XLSX_PATH, data_only=True, read_only=True)
    data = {name: records(book[name]) for name in book.sheetnames}
    cn = connect()
    cursor = cn.cursor()
    try:
        cursor.execute("SELECT Count(*) FROM [Товары]")
        if cursor.fetchone()[0]:
            raise RuntimeError("В таблице «Товары» уже есть записи. Импорт остановлен, чтобы не перезаписать вашу БД.")

        role_ids = {}
        for row in data["Роли"]:
            cursor.execute("INSERT INTO [Роли] ([Название]) VALUES (?)", row["Название"])
            cursor.execute("SELECT @@IDENTITY")
            role_ids[row["Название"]] = cursor.fetchone()[0]
        category_ids, manufacturer_ids, supplier_ids, warehouse_ids = {}, {}, {}, {}
        for sheet_name, table, target in [
            ("Категории", "Категории", category_ids),
            ("Производители", "Производители", manufacturer_ids),
            ("Поставщики", "Поставщики", supplier_ids),
        ]:
            for row in data[sheet_name]:
                cursor.execute(f"INSERT INTO [{table}] ([Название]) VALUES (?)", row["Название"])
                cursor.execute("SELECT @@IDENTITY")
                target[row["Название"]] = cursor.fetchone()[0]
        for row in data["Склады"]:
            cursor.execute("INSERT INTO [Склады] ([Название], [Адрес]) VALUES (?, ?)", row["Название"], row["Адрес"])
            cursor.execute("SELECT @@IDENTITY")
            warehouse_ids[row["Название"]] = cursor.fetchone()[0]
        for row in data["Пользователи"]:
            cursor.execute("INSERT INTO [Пользователи] ([КодРоли], [ФИО], [Логин], [Пароль]) VALUES (?, ?, ?, ?)",
                           role_ids[row["Роль"]], row["ФИО"], row["Логин"], password_hash(row["Пароль"]))

        product_ids = {}
        for row in data["Товары"]:
            cursor.execute("""INSERT INTO [Товары]
                ([Артикул], [Наименование], [КодКатегории], [Описание], [КодПроизводителя],
                 [КодПоставщика], [Цена], [Единица], [Скидка], [Фото])
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                row["Артикул"], row["Наименование"], category_ids[row["Категория"]], row["Описание"],
                manufacturer_ids[row["Производитель"]], supplier_ids[row["Поставщик"]], row["Цена"],
                row["Единица"], row["Скидка"], row["Фото"] or None)
            cursor.execute("SELECT @@IDENTITY")
            product_ids[row["Артикул"]] = cursor.fetchone()[0]
        for row in data["Остатки"]:
            cursor.execute("INSERT INTO [Остатки] ([КодТовара], [КодСклада], [Количество]) VALUES (?, ?, ?)",
                           product_ids[row["Артикул"]], warehouse_ids[row["Склад"]], row["Количество"])

        customer_ids = {}
        for row in data["Покупатели"]:
            cursor.execute("INSERT INTO [Покупатели] ([ФИО], [Телефон], [Адрес]) VALUES (?, ?, ?)",
                           row["ФИО"], row["Телефон"], row["Адрес"])
            cursor.execute("SELECT @@IDENTITY")
            customer_ids[row["ФИО"]] = cursor.fetchone()[0]
        for row in data["Пользователи"]:
            if row["Роль"] == "Авторизованный клиент" and row["ФИО"] in customer_ids:
                cursor.execute("UPDATE [Пользователи] SET [КодПокупателя]=? WHERE [Логин]=?",
                               customer_ids[row["ФИО"]], row["Логин"])
        order_ids = {}
        for row in data["Заказы"]:
            order_date = date.fromisoformat(str(row["ДатаЗаказа"]))
            delivery_date = date.fromisoformat(str(row["ДатаДоставки"]))
            cursor.execute("""INSERT INTO [Заказы]
                ([ДатаЗаказа], [ДатаДоставки], [КодПокупателя], [КодСклада], [АдресДоставки], [КодПолучения], [Статус])
                VALUES (?, ?, ?, ?, ?, ?, ?)""",
                order_date, delivery_date, customer_ids[row["Покупатель"]], warehouse_ids[row["Склад"]],
                row["Адрес"], row["Код"], row["Статус"])
            cursor.execute("SELECT @@IDENTITY")
            order_ids[row["Номер"]] = cursor.fetchone()[0]
        prices = {row["Артикул"]: float(row["Цена"]) * (100 - float(row["Скидка"])) / 100
                  for row in data["Товары"]}
        for row in data["СоставЗаказа"]:
            cursor.execute("INSERT INTO [СоставЗаказа] ([КодЗаказа], [КодТовара], [Количество], [ЦенаНаМоментЗаказа]) VALUES (?, ?, ?, ?)",
                           order_ids[row["Номер"]], product_ids[row["Артикул"]], row["Количество"], prices[row["Артикул"]])
        cn.commit()
        print(f"Импорт завершён: {len(product_ids)} товаров, {len(order_ids)} заказов.")
        print("Условные учётные записи: client@mebel.local / Client2026; manager@mebel.local / Manager2026; admin@mebel.local / Admin2026")
    except Exception:
        cn.rollback()
        raise
    finally:
        cn.close()
        book.close()


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"Ошибка импорта: {exc}", file=sys.stderr)
        raise SystemExit(1)
