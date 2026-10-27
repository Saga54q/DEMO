"""Импорт табличных ресурсов в SQLite. Источник — книга Furniture_demo_import.xlsx."""
from pathlib import Path
import shutil

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "import" / "Furniture_demo_import.xlsx"


def _sheet_rows(workbook, name):
    values = list(workbook[name].iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(value).strip() if value is not None else "" for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:]
            if row and any(value is not None and str(value).strip() for value in row)]


def _text(value):
    return "" if value is None else str(value).strip()


def _required(mapping, value, label):
    key = _text(value)
    if key not in mapping:
        raise ValueError(f"В исходных данных не найдено значение «{key}» для {label}.")
    return mapping[key]


def load_data(connection, source=SOURCE):
    """Импортирует согласованный учебный набор и корректно закрывает Excel-книгу."""
    if not Path(source).exists():
        raise FileNotFoundError(f"Не найден Excel-источник: {source}")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheets = {name: _sheet_rows(workbook, name) for name in workbook.sheetnames}
    finally:
        workbook.close()

    def insert_lookup(table, column, rows):
        values = list(dict.fromkeys(_text(row[column]) for row in rows if _text(row.get(column))))
        connection.executemany(
            f"INSERT INTO {table}(name) VALUES(?)", [(value,) for value in values])
        id_field = {"Role": "role_id", "Category": "category_id",
                    "Manufacturer": "manufacturer_id", "Supplier": "supplier_id",
                    "Unit": "unit_id", "OrderStatus": "status_id"}[table]
        return {row[0]: row[1] for row in connection.execute(
            f"SELECT name, {id_field} FROM {table}")}

    roles = insert_lookup("Role", "name", sheets["Roles"])
    categories = insert_lookup("Category", "name", sheets["Categories"])
    manufacturers = insert_lookup("Manufacturer", "name", sheets["Manufacturers"])
    suppliers = insert_lookup("Supplier", "name", sheets["Suppliers"])
    units = insert_lookup("Unit", "name", sheets["Units"])
    statuses = insert_lookup("OrderStatus", "name", sheets["OrderStatuses"])

    warehouses = {}
    for row in sheets["Warehouses"]:
        name = _text(row.get("name"))
        address = _text(row.get("address"))
        if name:
            cursor = connection.execute(
                "INSERT INTO Warehouse(name,address) VALUES(?,?)", (name, address))
            warehouses[name] = cursor.lastrowid

    users = []
    for row in sheets["Users"]:
        users.append((roles[_text(row["role"])], _text(row["fio"]),
                      _text(row["login"]), _text(row["password"])))
    connection.executemany(
        "INSERT INTO User(role_id,fio,login,password) VALUES(?,?,?,?)", users)

    product_ids = {}
    for row in sheets["Products"]:
        article = _text(row["article"])
        image_source = ROOT / "import" / _text(row.get("image_path")) if _text(row.get("image_path")) else None
        image_path = None
        if image_source and image_source.is_file():
            target = ROOT / "images" / image_source.name
            target.parent.mkdir(parents=True, exist_ok=True)
            if image_source.resolve() != target.resolve():
                shutil.copy2(image_source, target)
            image_path = f"images/{target.name}"
        cursor = connection.execute("""
            INSERT INTO Product(article,name,category_id,description,manufacturer_id,
                supplier_id,price,unit_id,discount,image_path)
            VALUES(?,?,?,?,?,?,?,?,?,?)
        """, (article, _text(row["name"]), _required(categories, row["category"], "категория"),
              _text(row.get("description")), _required(manufacturers, row["manufacturer"], "производитель"),
              _required(suppliers, row["supplier"], "поставщик"), float(row["price"] or 0),
              _required(units, row["unit"], "единица измерения"), float(row["discount"] or 0), image_path))
        product_ids[article] = cursor.lastrowid

    for row in sheets["Inventory"]:
        connection.execute("""
            INSERT INTO Inventory(product_id,warehouse_id,quantity) VALUES(?,?,?)
        """, (product_ids[_text(row["article"])], warehouses[_text(row["warehouse"])],
              int(row["quantity"] or 0)))

    customers = {}
    for row in sheets["Customers"]:
        fio = _text(row["fio"])
        cursor = connection.execute(
            "INSERT INTO Customer(fio,phone,address) VALUES(?,?,?)",
            (fio, _text(row.get("phone")), _text(row.get("address"))))
        customers[fio] = cursor.lastrowid

    order_ids = {}
    for row in sheets["Orders"]:
        order_id = int(row["order_id"])
        cursor = connection.execute("""
            INSERT INTO ShopOrder(order_date,delivery_date,customer_id,warehouse_id,
                delivery_address,code,status_id)
            VALUES(?,?,?,?,?,?,?)
        """, (_text(row["order_date"]), _text(row["delivery_date"]),
              _required(customers, row["customer_fio"], "клиент"),
              _required(warehouses, row["warehouse"], "склад"),
              _text(row.get("delivery_address")), _text(row.get("code")),
              _required(statuses, row["status"], "статус заказа")))
        order_ids[order_id] = cursor.lastrowid

    for row in sheets["OrderItems"]:
        product_id = product_ids[_text(row["article"])]
        price = connection.execute("SELECT price FROM Product WHERE product_id=?", (product_id,)).fetchone()[0]
        connection.execute("""
            INSERT INTO OrderItem(order_id,product_id,quantity,unit_price) VALUES(?,?,?,?)
        """, (order_ids[int(row["order_id"])], product_id, int(row["quantity"]), price))

    counts = {table: connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
              for table in ("Product", "User", "Customer", "Warehouse", "ShopOrder", "OrderItem")}
    print("Импорт учебных данных завершён:", counts)
    return counts
