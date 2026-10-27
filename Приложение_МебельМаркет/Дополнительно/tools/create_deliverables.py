from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
FONT_REGULAR = FONT_DIR / "DejaVuSans.ttf"
FONT_BOLD = FONT_DIR / "DejaVuSans-Bold.ttf"


def font(size, bold=False):
    path = FONT_BOLD if bold else FONT_REGULAR
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()


def make_erd():
    image = Image.new("RGB", (1800, 1040), "#f3f7f6")
    draw = ImageDraw.Draw(image)
    title_font = font(32, True)
    header_font = font(19, True)
    body_font = font(15)
    small_font = font(14, True)
    draw.text((55, 28), "МЕБЕЛЬМАРКЕТ · ER-диаграмма базы данных", font=title_font, fill="#124b49")
    nodes = {
        "Роли": (40, 100, ["КодРоли · PK", "Название · уникальное"]),
        "Пользователи": (390, 100, ["КодПользователя · PK", "КодРоли · FK", "КодПокупателя · FK, null", "Логин · уникальный"]),
        "Покупатели": (720, 100, ["КодПокупателя · PK", "ФИО · телефон · адрес"]),
        "Заказы": (1075, 100, ["КодЗаказа · PK", "КодПокупателя · FK", "КодСклада · FK", "Даты · статус · адрес"]),
        "Категории": (30, 400, ["КодКатегории · PK", "Название · уникальное"]),
        "Производители": (325, 400, ["КодПроизводителя · PK", "Название · уникальное"]),
        "Поставщики": (620, 400, ["КодПоставщика · PK", "Название · уникальное"]),
        "Товары": (915, 400, ["КодТовара · PK", "Артикул · уникальный", "КодКатегории · FK", "КодПроизводителя · FK", "КодПоставщика · FK", "Цена · скидка · фото"]),
        "СоставЗаказа": (1285, 400, ["КодЗаказа · PK, FK", "КодТовара · PK, FK", "Количество", "ЦенаНаМоментЗаказа"]),
        "Склады": (1465, 770, ["КодСклада · PK", "Название · адрес"]),
        "Остатки": (845, 770, ["КодТовара · PK, FK", "КодСклада · PK, FK", "Количество"]),
    }
    w, h = 265, 185
    centers = {name: (x + w // 2, y + h // 2) for name, (x, y, _) in nodes.items()}
    def label(text, x, y):
        tw = draw.textbbox((0, 0), text, font=small_font)[2]
        draw.rounded_rectangle((x-tw//2-6, y-12, x+tw//2+6, y+12), radius=5, fill="#f3f7f6")
        draw.text((x-tw//2, y-10), text, font=small_font, fill="#276d66")
    def line(points, text=None, label_at=None):
        draw.line(points, fill="#84aaa3", width=3)
        if text and label_at:
            label(text, *label_at)
    # Горизонтальные связи верхнего ряда.
    line([(305, 192), (390, 192)], "1 : N", (347, 166))
    line([(720, 205), (655, 205)], "0..1 : N", (687, 180))
    line([(985, 192), (1075, 192)], "1 : N", (1030, 166))
    # Справочники идут отдельными ортогональными маршрутами ниже карточек.
    for x_start, route_y, x_end, text in [(162, 635, 960, "1 : N"), (457, 665, 1010, "1 : N"), (752, 695, 1060, "1 : N")]:
        line([(x_start, 585), (x_start, route_y), (x_end, route_y), (x_end, 585)], text, ((x_start+x_end)//2, route_y-14))
    # Остатки и строки заказа — связующие сущности.
    line([(1180, 535), (1285, 535)], "1 : N", (1232, 509))
    line([(1200, 285), (1417, 400)], "1 : N", (1302, 326))
    line([(1045, 585), (978, 770)], "1 : N", (1010, 676))
    line([(1110, 862), (1465, 862)], "1 : N", (1288, 836))
    # Склад связан с заказом через внешний маршрут по правому краю схемы.
    line([(1730, 862), (1770, 862), (1770, 192), (1340, 192)], "1 : N", (1750, 530))
    for name, (x, y, fields) in nodes.items():
        draw.rounded_rectangle((x, y, x+w, y+h), radius=15, fill="#ffffff", outline="#19837a", width=3)
        draw.rounded_rectangle((x, y, x+w, y+40), radius=13, fill="#087f78")
        draw.rectangle((x, y+28, x+w, y+42), fill="#087f78")
        draw.text((x+12, y+8), name, font=header_font, fill="white")
        for idx, text in enumerate(fields):
            draw.text((x+12, y+51+idx*22), text, font=body_font, fill="#263a36")
    draw.text((55, 925), "PK — первичный ключ · FK — внешний ключ · таблицы справочников и связей показаны отдельными сущностями", font=font(16), fill="#526863")
    image.save(ROOT / "Диаграмма_связей.png", dpi=(160, 160))


def shade(cell, color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def make_report():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(38, 54, 50)
    for style_name, size, color in [("Title", 28, "087F78"), ("Heading 1", 18, "087F78"), ("Heading 2", 13, "124B49")]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ОТЧЁТ\n")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(90, 110, 104)
    p = doc.add_paragraph("Разработка учебного приложения\n«МебельМаркет»", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Реляционная база данных Microsoft Access · приложение Python / Tkinter")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("\nПроектный отчёт и описание реализации", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("\nИсходные приложения экзамена в комплекте отсутствуют. Сведения, визуальный стиль и набор товаров в демонстрационной базе являются условными.").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("1. Назначение и границы проекта", level=1)
    doc.add_paragraph("Цель — подготовить минимально необходимую нормализованную структуру базы мебельного магазина на Access, нативный стандартный модуль DAO/VBA для создания объектов и простое Python-приложение для работы с каталогом, остатками и заказами. Интерфейс Tkinter запускается в Windows и взаимодействует с .accdb через pyodbc и Microsoft Access ODBC Driver.")
    doc.add_paragraph("Фотографии основных условий сохранены в папке `Фото_условия` и использованы для требований к ролям, карточке товара и алгоритму. Отдельные приложения `Прил_1_ОЗ_КОД 09.02.07-2-2026-М1.docx`, `Прил_2_ОЗ_КОД 09.02.07-2-2026-М1.rar` и `Прил_3_ОЗ_КОД 09.02.07-2-2026-М2.docx` не предоставлены: поэтому тестовые записи, логотип, заглушка и оформление за пределами прямых правил условия являются условными. Полный текст отдельного пункта 7 из запроса на подготовку комплекта не поступил.")
    doc.add_heading("2. Функциональные требования", level=1)
    for item in [
        "первое окно — вход с логином/паролем из БД либо переход к гостевому каталогу; выход из учётной записи возвращает к окну входа, в шапке отображается ФИО;",
        "гость и авторизованный клиент просматривают товары без фильтрации, сортировки и поиска; менеджер и администратор получают поиск и сортировку; менеджер просматривает заказы;",
        "администратор может добавлять, редактировать и удалять товары и заказы;",
        "карточка показывает фото/заглушку, категорию, название, описание, производителя, поставщика, цену, единицу, остатки по складам и скидку;",
        "скидка строго больше 15% выделяется фоном #008080; товар без остатка — серым фоном; прежняя цена зачёркнута и красная, цена со скидкой указана рядом;",
        "дополнительно: корзина клиента, проверка наличия на складе, заказ, история заказов и транзакционное уменьшение остатков."
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("3. Состав и модель данных", level=1)
    doc.add_paragraph("БД включает 11 таблиц. Справочники вынесены отдельно, а отношения многие-ко-многим между заказами и товарами реализованы через СоставЗаказа. Остатки отражаются отдельно для каждой пары товар–склад.")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ["Таблица", "Назначение"]):
        cell.text = text
        shade(cell, "087F78")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    descriptions = [
        ("Роли", "Роли доступа."), ("Пользователи", "Учётные записи, роль и связь с профилем покупателя."),
        ("Категории", "Классификация каталога."), ("Производители", "Справочник изготовителей."),
        ("Поставщики", "Справочник поставщиков."), ("Склады", "Места хранения и адреса."),
        ("Товары", "Артикул, наименование, цена, скидка и внешние ключи справочников."),
        ("Остатки", "Количество каждого товара в каждом складе."),
        ("Покупатели", "ФИО и контактные данные."), ("Заказы", "Покупатель, склад, даты, код и состояние."),
        ("СоставЗаказа", "Товары, количество и зафиксированная цена на момент продажи.")]
    for name, description in descriptions:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = description
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph("Ключевые ограничения: уникальность логинов и артикулов; уникальные названия справочников; составные первичные ключи в Остатках и СоставЗаказа; неотрицательные цена/остаток; скидка 0–100%; положительное количество заказа; внешние ключи.")
    doc.add_heading("4. Диаграмма сущностей и связей", level=1)
    doc.add_paragraph("Схема фиксирует отношения справочников, покупателей, складских остатков и строк заказов. PK обозначает первичный ключ, FK — внешний ключ.")
    doc.add_picture(str(ROOT / "Диаграмма_связей.png"), width=Cm(16.8))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("5. Реализация и запуск", level=1)
    doc.add_paragraph("Структуру базы создаёт импортируемый стандартный модуль Access VBA `Модуль_создания_БД.bas` с процедурой `BuildFurnitureDatabase`. Он создаёт таблицы, индексы, проверки, DAO-связи и четыре сохранённых запроса. Пользователь создаёт пустой .accdb, импортирует BAS и запускает макрос в Microsoft Access.")
    doc.add_paragraph("Python-приложение находится в `Приложение/main.py`; импортёр демонстрационных данных — в `Приложение/ЗаполнитьБД.py`. Файл Excel формируется скриптом `БД/СоздатьДемоДанные.py`. Две отдельные инструкции описывают запуск в PyCharm и VS Code. Блок-схема алгоритма подготовлена в PDF по обозначениям ГОСТ 19.701-90. Бинарный .accdb не включён: он должен быть создан локально в Access.")
    doc.add_heading("6. Проверки и ограничения", level=1)
    doc.add_paragraph("В среде подготовки выполнена синтаксическая компиляция Python-файлов (`py_compile`), сформирован файл условных данных Excel. Среда не содержит Windows, Microsoft Access, Access ODBC Driver и полноценно проверяемого Tk GUI. Поэтому проверка выполнения VBA, создания реальной .accdb, ODBC-запросов и пользовательских сценариев на Access не проведена.")
    doc.add_paragraph("Перед демонстрацией требуется проверить модуль на целевой версии Access, создать новую базу, выполнить импорт и пройти роли по условию: гость, клиент, менеджер, администратор; проверить показ карточек, подсветку, поиск/сортировку сотрудника и CRUD администратора. Отдельно можно проверить дополнительный сценарий заказа и списания остатка. При несовместимости конкретной версии ACE/DAO может потребоваться корректировка SQL и повторный тест.")
    doc.add_paragraph("SHA-256 в учебном примере не заменяет промышленное хранение паролей. Для настоящего сервиса нужны соль, медленный алгоритм хеширования, серверная БД, авторизация на сервере и аудит доступа.")
    doc.add_heading("7. Содержание поставки", level=1)
    doc.add_paragraph("Структура проекта и полный список материалов описаны в `Дополнительно/Руководство.md`. Краткий план демонстрации и ответы по модели данных находятся в `Шпаргалка_экзамен.md`. Скрипты диаграмм и документов находятся в `Дополнительно/tools/`.")
    doc.add_paragraph("Официальные приложения задания, подлинные скриншоты и подтверждение запуска на Windows не добавлены и не имитируются.")
    out = ROOT / "Отчет_МебельМаркет.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    make_erd()
    make_report()
