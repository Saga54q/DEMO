"""Одноразовая загрузка предоставленных Excel-файлов в SQLite."""
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "import"
IMAGES = ROOT / "images"


def rows(book, sheet):
    """Читает строки листа, используя первую строку как имена столбцов."""
    ws = load_workbook(SOURCE / book, data_only=True, read_only=True)[sheet]
    values = list(ws.iter_rows(values_only=True))
    headers = values[0]
    return [dict(zip(headers, row)) for row in values[1:]
            if row and row[0] is not None]


def iso(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value)


def load_excel_data(connection):
    """Импортирует справочники до товаров и заказов, чтобы не нарушать FK."""
    products = rows("Tovar.xlsx", "Product")
    suppliers = rows("Tovar.xlsx", "Supplier")
    makers = rows("Tovar.xlsx", "Manufacturer")
    types = rows("Tovar.xlsx", "Category")
    users = rows("User_import.xlsx", "User")
    roles = rows("User_import.xlsx", "Role")
    order_file = "Order_import1.xlsx"
    customers = rows(order_file, "Fio")
    statuses = rows(order_file, "Status")
    points = rows("PickUpPoint_import1.xlsx", "PickUpPoint")
    orders = rows(order_file, "Order")
    items = rows(order_file, "Count")

    connection.executemany("INSERT INTO Role VALUES(?,?)",
                           [(r["RoleId"], r["RoleName"]) for r in roles])
    connection.executemany("INSERT INTO User VALUES(?,?,?,?,?)",
                           [(u["User_Id"], u["RoleId"], u["FIO"], str(u["Login"]),
                             str(u["Password"])) for u in users])
    connection.executemany("INSERT INTO Supplier VALUES(?,?)",
                           [(r["SupplierId"], r["SupplierName"].strip()) for r in suppliers])
    connection.executemany("INSERT INTO Manufacturer VALUES(?,?)",
                           [(r["ManufacturerId"], r["ManufacturerName"].strip()) for r in makers])
    connection.executemany("INSERT INTO ProductType VALUES(?,?)",
                           [(r["CategoryId"], r["CategoryName"].strip()) for r in types])
    connection.executemany("INSERT INTO Customer VALUES(?,?)",
                           [(r["Id"], r["Fio"].strip()) for r in customers])
    connection.executemany("INSERT INTO OrderStatus VALUES(?,?)",
                           [(r["StatusId"], r["StatusName"]) for r in statuses])
    connection.executemany("INSERT INTO PickupPoint VALUES(?,?,?,?,?)",
                           [(p["Id"], str(p["CodePoint"]), p["City"].strip(),
                             p["Street"].strip(), str(p["Number"])) for p in points])

    IMAGES.mkdir(exist_ok=True)
    product_rows = []
    for p in products:
        image = p.get("Photo")
        image_path = IMAGES / str(image) if image else None
        source_image = SOURCE / str(image) if image else None
        if image_path and source_image and source_image.exists():
            with Image.open(source_image) as photo:
                photo.convert("RGB").save(image_path, "JPEG", quality=90)
        stored_image = f"images/{image}" if image_path and image_path.exists() else None
        product_rows.append((p["Id"], p["Article"], p["ProductName"], p["Edenica Izmerenia"],
            float(p["Price"] or 0), p["SupplierId"], p["ManufacturerId"], p["CategoryId"],
            float(p["Skidka"] or 0), int(p["Sklad"] or 0),
            p.get("Description") or "", stored_image))
    connection.executemany("INSERT INTO Product VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", product_rows)

    connection.executemany("INSERT INTO ShopOrder VALUES(?,?,?,?,?,?,?)", [
        (o["Order_Id"], iso(o["Date"]), iso(o["Delivery"]), o["PickupPointId"],
         o["Fio_Id"], str(o["Code"]), o["Status"]) for o in orders])
    connection.executemany("INSERT INTO OrderItem VALUES(?,?,?)", [
        (i["Id"], i["Name"], int(i["Count"])) for i in items])
