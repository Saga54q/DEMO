"""Одноразовая очистка и загрузка исходных Excel-файлов Варианта 2."""
from datetime import date, datetime
from pathlib import Path
import re
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "import"
IMAGES = ROOT / "images"


def read_rows(filename):
    workbook = load_workbook(SOURCE / filename, data_only=True, read_only=True)
    try:
        values = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = values[0]
    return [dict(zip(headers, row)) for row in values[1:]
            if row and any(value is not None for value in row)]


def clean(value):
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip()


def date_text(value, order_id):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = clean(value)
    try:
        return datetime.strptime(text, "%d.%m.%Y").date().isoformat()
    except ValueError:
        if order_id == 7 and text == "30.02.2023":
            # В исходном Excel дата невозможна; нормализована к последнему дню февраля.
            print("Исправлено: заказ 7, 30.02.2023 → 28.02.2023")
            return "2023-02-28"
        raise ValueError(f"Некорректная дата заказа №{order_id}: {text}")


def unique_ids(values):
    names = list(dict.fromkeys(clean(value) for value in values if clean(value)))
    return {name: index for index, name in enumerate(names, 1)}


def load_excel_data(connection):
    products = read_rows("Tovar.xlsx")
    users = read_rows("user_import.xlsx")
    orders = read_rows("orders.xlsx")

    role_ids = unique_ids(row["Роль сотрудника"] for row in users)
    supplier_ids = unique_ids(row["Поставщик"] for row in products)
    maker_ids = unique_ids(row["Производитель"] for row in products)
    category_ids = unique_ids(row["Категория товара"] for row in products)
    customer_ids = unique_ids(row["ФИО авторизированного клиента"] for row in orders)
    status_ids = unique_ids(clean(row["Статус заказа"]) for row in orders)

    connection.executemany("INSERT INTO Role VALUES(?,?)",
                           [(number, name) for name, number in role_ids.items()])
    connection.executemany("INSERT INTO Supplier VALUES(?,?)",
                           [(number, name) for name, number in supplier_ids.items()])
    connection.executemany("INSERT INTO Manufacturer VALUES(?,?)",
                           [(number, name) for name, number in maker_ids.items()])
    connection.executemany("INSERT INTO ProductType VALUES(?,?)",
                           [(number, name) for name, number in category_ids.items()])
    connection.executemany("INSERT INTO Customer VALUES(?,?)",
                           [(number, name) for name, number in customer_ids.items()])
    connection.executemany("INSERT INTO OrderStatus VALUES(?,?)",
                           [(number, name) for name, number in status_ids.items()])

    connection.executemany("INSERT INTO User VALUES(?,?,?,?,?)", [
        (number, role_ids[clean(row["Роль сотрудника"])], clean(row["ФИО"]),
         clean(row["Логин"]), clean(row["Пароль"]))
        for number, row in enumerate(users, 1)])

    # В файле пункты выдачи заданы строками без заголовка; их ID — номер строки.
    point_book = load_workbook(SOURCE / "pickup_points.xlsx", data_only=True, read_only=True)
    try:
        point_rows = [row[0] for row in point_book.active.iter_rows(values_only=True)
                      if row and row[0]]
    finally:
        point_book.close()
    points = []
    for point_id, address in enumerate(point_rows, 1):
        parts = [clean(part) for part in clean(address).split(",")]
        postal = parts[0] if parts else ""
        city = parts[1] if len(parts) > 1 else ""
        street = parts[2] if len(parts) > 2 else ""
        number = parts[3] if len(parts) > 3 else ""
        points.append((point_id, postal, city, street, number))
    connection.executemany("INSERT INTO PickupPoint VALUES(?,?,?,?,?)", points)

    IMAGES.mkdir(exist_ok=True)
    product_ids = {}
    product_rows = []
    for product_id, row in enumerate(products, 1):
        article = clean(row["Артикул"])
        product_ids[article] = product_id
        photo = clean(row.get("Фото"))
        source_image = SOURCE / photo if photo else None
        image_path = IMAGES / photo if photo else None
        if source_image and source_image.exists():
            with Image.open(source_image) as image:
                image.convert("RGB").save(image_path, "JPEG", quality=90)
        image_name = f"images/{photo}" if image_path and image_path.exists() else None
        product_rows.append((product_id, article, clean(row["Наименование товара"]),
            clean(row["Единица измерения"]), float(row["Цена"]),
            supplier_ids[clean(row["Поставщик"])], maker_ids[clean(row["Производитель"])],
            category_ids[clean(row["Категория товара"])], float(row["Действующая скидка"] or 0),
            int(row["Кол-во на складе"] or 0), clean(row["Описание товара"]), image_name))
    connection.executemany("INSERT INTO Product VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", product_rows)

    pickup_ids = {point_id: point_id for point_id in range(1, len(points) + 1)}
    order_rows, item_rows = [], []
    for row in orders:
        order_id = int(row["Номер заказа"])
        status = clean(row["Статус заказа"])
        customer = clean(row["ФИО авторизированного клиента"])
        order_rows.append((order_id, date_text(row["Дата заказа"], order_id),
            date_text(row["Дата доставки"], order_id), pickup_ids[int(row["Адрес пункта выдачи"])],
            customer_ids[customer], clean(row["Код для получения"]), status_ids[status]))
        tokens = [part.strip() for part in clean(row["Артикул заказа"]).split(",")]
        if len(tokens) % 2:
            raise ValueError(f"В заказе №{order_id} нечетное число элементов в перечне товаров")
        for article, quantity in zip(tokens[::2], tokens[1::2]):
            if article not in product_ids:
                raise ValueError(f"В заказе №{order_id} нет артикула {article} в таблице товаров")
            item_rows.append((order_id, product_ids[article], int(quantity)))

    connection.executemany("INSERT INTO ShopOrder VALUES(?,?,?,?,?,?,?)", order_rows)
    connection.executemany("INSERT INTO OrderItem VALUES(?,?,?)", item_rows)
